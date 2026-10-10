"""분석 결과 Excel/CSV 내보내기."""
import csv
import datetime
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.analysis import EQUIP_NOTE

COLUMNS = ['위치', '아이템명', '아이템ID', '수량', '제련', '카드', '옵션',
           '최저가', '중앙가', '평균가', '노점수', '예상금액', '시세조회시간']
WIDTHS = [9, 34, 10, 8, 6, 30, 24, 13, 13, 13, 8, 16, 17]
MONEY_COLS = {'최저가', '중앙가', '평균가', '예상금액', '수량'}


def default_filename(folder, ext='xlsx'):
    return os.path.join(folder, '라그나로크_창고분석_%s.%s' % (datetime.datetime.now().strftime('%Y%m%d_%H%M'), ext))


def row_values(r, fmt):
    """fmt: card_text(entry), option_text(entry) 함수를 가진 객체(ui formatter)."""
    e, p = r.entry, r.price
    t = datetime.datetime.fromtimestamp(p.updated_at).strftime('%Y-%m-%d %H:%M') if p and p.updated_at else ''
    return [e.source, e.name, e.item_id, e.quantity, e.refine if e.refine else None,
            fmt.card_text(e), fmt.option_text(e),
            p.lowest_price if p else None, p.median_price if p else None, p.average_price if p else None,
            p.listing_count if p else None, r.value, t]


def export(path, rows, fmt, summary=None, meta=None):
    if path.lower().endswith('.csv'):
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            w.writerow(COLUMNS)
            for r in rows:
                w.writerow(row_values(r, fmt))
        return path

    wb = Workbook()
    ws = wb.active
    ws.title = '아이템 목록'
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append(row_values(r, fmt))
    for i, (name, w) in enumerate(zip(COLUMNS, WIDTHS), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        if name in MONEY_COLS:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = '#,##0'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet('요약', 0)
    for k, v in (meta or {}).items():
        s.append([k, v])
    if summary:
        s.append([])
        s.append(['전체 아이템 종류', summary['kinds']])
        s.append(['전체 아이템 수량', summary['quantity']])
        s.append(['시세 있는 종류', summary['priced_kinds']])
        s.append(['예상 총 판매금액', summary['total_value']])
        for k, v in summary['assets'].items():
            s.append(['%s 예상 자산' % k, v])
    s.append([])
    s.append(['주의', EQUIP_NOTE])
    for (a, b) in s.iter_rows(min_row=1, max_col=2):
        if isinstance(b.value, int):
            b.number_format = '#,##0'
    s.column_dimensions['A'].width, s.column_dimensions['B'].width = 22, 60
    wb.save(path)
    return path
