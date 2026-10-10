"""Replay 분석 결과 Excel / JSON 내보내기."""
import datetime
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill('solid', fgColor='305496')


def default_name(folder, prefix, ext):
    return os.path.join(folder, '%s_%s.%s' % (prefix, datetime.datetime.now().strftime('%Y%m%d_%H%M'), ext))


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), HEAD
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, len(str(h)) * 2 + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                c.number_format = '#,##0' if isinstance(c.value, int) or abs(c.value) >= 100 else '#,##0.00'
    ws.freeze_panes = 'A2'
    return ws


def _t(ms):
    s = int((ms or 0) // 1000)
    return '%02d:%02d:%02d' % (s // 3600, s // 60 % 60, s % 60)


def export_report_xlsx(path, rep, timeline_rows=None):
    wb = Workbook()
    wb.remove(wb.active)
    m, w, cb, k = rep['meta'], rep['window'], rep['combat'], rep['kpi']
    summary = [('캐릭터', m['character']), ('직업', m['job']), ('Level', m['level']), ('서버', m['server']),
               ('Replay', m['replay_file']), ('녹화 시작', m['start_time']), ('Replay 길이', _t(m['duration_ms'])),
               ('맵', ', '.join(m.get('map_names') or m['maps'])), ('분석 구간', w['label']),
               ('구간 길이', _t(w['duration_ms'])), ('', '')]
    summary += [('● 총 피해', cb['total_damage'] if cb['available'] else None), ('○ Combat DPS', k['combat_dps']),
                ('○ 전체 DPS', k['total_dps']), ('● 처치', k['kills']), ('○ Base EXP/h', k['base_exp_per_hour']),
                ('○ Zeny 순변화', k['zeny_net']), ('○ 드랍 가치', k['loot_value']),
                ('△ 소모품 비용', k['consumable_cost']), ('○ 순이익', k['profit']), ('○ Profit/h', k['profit_per_hour']),
                ('', '')] + [('요약', line) for line in rep['summary_text']]
    _sheet(wb, 'Summary', ['항목', '값'], summary, {'항목': 18, '값': 70})

    _sheet(wb, 'Combat', ['항목', '값'], [
        (key, cb.get(key)) for key in ('total_damage', 'hit_count', 'attack_count', 'normal_attacks', 'skill_hits',
                                       'miss', 'critical', 'avg_hit', 'max_hit', 'min_hit', 'combat_ms', 'dps_total',
                                       'dps_combat', 'kills', 'kill_time_avg', 'kill_time_min', 'kill_time_max',
                                       'taken_total', 'taken_avg', 'taken_max', 'hp_recovered')]
        + [('burst_%ss_dps' % b, v['dps']) for b, v in cb['burst'].items()] + [('note', cb['note'])],
        {'항목': 20, '값': 20})
    _sheet(wb, 'Skills', ['스킬', 'ID', '사용', '캐스팅', '명중', '총피해', '평균', '최소', '최대', '기여율%', 'DPS',
                          '평균간격(s)', '최소간격(s)', '평균캐스팅(ms)', '관측효율%', '기대딜레이효율%'],
           [(s['name'], s['skill_id'], s['uses'], s['casts'], s['hits'], s['total_damage'], s['avg_damage'],
             s['min_damage'], s['max_damage'], s['share'], s['dps'], s['avg_interval'], s['min_interval'],
             s['cast_ms_avg'], s['efficiency_observed'], s['efficiency_expected']) for s in rep['skills']['skills']],
           {'스킬': 24})
    _sheet(wb, 'Buffs', ['버프', 'EFST', '가동률%', '활성(s)', '공백합(s)', '재사용', '공백 구간', 'DPS(버프O)', 'DPS(버프X)'],
           [(b['name'], b['efst'], b['uptime'], b['active_ms'] / 1000, b['gap_total_ms'] / 1000, b['recasts'],
             ', '.join('%s~%s' % (_t(g['start_ms']), _t(g['end_ms'])) for g in b['gaps'][:20]),
             b['dps_on'], b['dps_off']) for b in rep['buffs']['buffs']], {'버프': 22, '공백 구간': 50})
    _sheet(wb, 'Monsters', ['몬스터', 'ID', '조우', '처치', '총피해', '평균피해', '평균Kill(s)', '최소', '최대', '받은피해'],
           [(r['name'], r['class_id'], r['encounters'], r['kills'], r['damage'], r['avg_damage'], r['kill_time_avg'],
             r['kill_time_min'], r['kill_time_max'], r['taken']) for r in rep['monsters']['monsters']], {'몬스터': 24})
    lt = rep['loot']
    _sheet(wb, 'Loot', ['구분', '아이템', 'ID', '수량', '단가', '가치', '노점수'],
           [('획득', r['name'], r['item_id'], r['quantity'], r['unit_price'], r['value'], r['listings'])
            for r in lt.get('loot', [])] +
           [('소모(△)', r['name'], r['item_id'], r['quantity'], r['unit_price'], r['value'], r['listings'])
            for r in lt.get('consumables', [])] + [('방법', lt.get('method') or lt.get('note'), '', '', '', '', '')],
           {'아이템': 30})
    ex = rep['exp']
    _sheet(wb, 'EXP', ['항목', '값'], [('Base 총', ex.get('base_total')), ('Job 총', ex.get('job_total')),
                                       ('Base/h', ex.get('base_per_hour')), ('Job/h', ex.get('job_per_hour')),
                                       ('비고', ex.get('note'))] +
           [('맵 %s Base' % mp, v['base']) for mp, v in (ex.get('by_map') or {}).items()], {'값': 40})
    if timeline_rows is not None:
        _sheet(wb, 'Timeline', ['시간', '분류', '내용', '신뢰도'],
               [(_t(t), c, d, conf) for t, c, d, conf, _ in timeline_rows[:200_000]], {'내용': 60})
    wb.save(path)
    return path


def export_compare_xlsx(path, cmp):
    wb = Workbook()
    wb.remove(wb.active)
    labels = cmp['labels']
    _sheet(wb, 'Compare', ['항목'] + labels + ['차이%% (%s 대비)' % labels[0]] * (len(labels) - 1),
           [[r['metric']] + r['values'] + r['diffs'][1:] for r in cmp['rows']] +
           [['경고: ' + w] for w in cmp['warnings']], {'항목': 20})
    _sheet(wb, 'Skill Damage', ['스킬'] + labels + ['차이%'] * (len(labels) - 1),
           [[s['name']] + s['damage'] + s['diffs'][1:] for s in cmp['skill_rows']], {'스킬': 24})
    _sheet(wb, 'Buff Uptime', ['버프'] + labels, [[b['name']] + b['uptime'] for b in cmp['buff_rows']], {'버프': 24})
    wb.save(path)
    return path


def export_json(path, obj):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=1, default=str)
    return path
